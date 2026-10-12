from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = (
    "用例编号", "用例名称", "用例类别", "是否运行", "最终状态",
    "执行耗时（秒）", "大致步骤", "验收标准", "结论说明",
)
STATUS_LABELS = {
    "PASSED": "通过", "FAILED": "未通过", "BLOCKED": "前置阻塞", "NOT_RUN": "未执行",
}
GENERIC_REASONS = {
    "当前计划版本首轮执行通过", "当前计划版本的首轮执行未通过", "尚无首轮执行记录",
}


def _display_text(value: object) -> str:
    text = str(value or "")
    for pattern, replacement in (
        (r"\bAndroid\b", "安卓手机"), (r"\biOS\b", "苹果手机"),
        (r"\bApp\b", "应用"), (r"\bLive\b", "动态照片"),
        (r"\bPOI\b", "地点信息"), (r"\bUI\b", "页面"),
        (r"\bID\b", "编号"), (r"\bTab\b", "页签"),
        (r"\bkcal\b", "千卡"), (r"\bAI\b", "智能"),
        (r"\bPro\b", "专业版"),
        (r"(\d+)mg\b", r"\1 毫克"), (r"(\d+)ml\b", r"\1 毫升"),
        (r"(\d+)g\b", r"\1 克"),
    ):
        text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)
    return text.replace("---", "：")


def _run_result(runs_dir: Path, run_id: object) -> dict | None:
    if not run_id:
        return None
    root = runs_dir.resolve()
    result_path = (root / str(run_id) / "result.json").resolve()
    if not result_path.is_relative_to(root) or not result_path.is_file():
        return None
    try:
        result = json.loads(result_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    return result if isinstance(result, dict) else None


def _elapsed_seconds(result: dict | None) -> float | None:
    if not result or not result.get("started_at") or not result.get("finished_at"):
        return None
    try:
        started = datetime.fromisoformat(result["started_at"])
        finished = datetime.fromisoformat(result["finished_at"])
        seconds = (finished - started).total_seconds()
    except (TypeError, ValueError):
        return None
    return round(seconds, 1) if seconds >= 0 else None


def _observation(result: dict | None, *, failed: bool) -> str:
    for oracle in (result or {}).get("oracles") or []:
        if not isinstance(oracle, dict) or bool(oracle.get("passed")) == failed:
            continue
        message = str(oracle.get("message") or "")
        if "实际观察：" in message:
            return _display_text(message.split("实际观察：", 1)[1].split("；依据：", 1)[0])[:200]
    return ""


def _conclusion(row: dict, result: dict | None) -> str:
    status = row["status"]
    reason = str(row.get("reason") or "")
    if status == "NOT_RUN" and reason == "尚无首轮执行记录":
        return "尚未执行，暂无验收结论。"
    if reason and reason not in GENERIC_REASONS:
        return _display_text(reason)
    expected = _display_text(row.get("expected"))
    if status == "PASSED":
        observed = _observation(result, failed=False)
        if observed:
            return f"通过：实际看到{observed}，符合验收要求。"
        return f"通过：按步骤完成，并核对了“{expected}”，结果符合要求。" if expected else "通过：步骤和验收项均已通过。"
    if status == "FAILED":
        observed = _observation(result, failed=True)
        if observed:
            return f"未通过：实际看到{observed}，未达到验收要求。"
        return f"未通过：未满足“{expected}”。" if expected else "未通过：本次验收未满足要求。"
    return _display_text(reason or "前置条件不足，尚不能完成验收。")


def _text_cell(sheet, row: int, column: int, value: str) -> None:
    cell = sheet.cell(row, column, value)
    cell.data_type = "s"  # Imported case text must never become an Excel formula.


def build_execution_workbook(rows: list[dict], runs_dir: Path) -> bytes:
    """Export the dashboard's effective case outcomes, not raw run pass/fail flags."""
    rows = [{**row, "qualification_run_id": row.get("qualification_run_id") or row.get("first_run_evidence_run_id")} for row in rows]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用例结果"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A10"

    sheet["A2"] = "用例执行结果"
    sheet["A2"].font = Font(name="微软雅黑", size=15, bold=True, color="152C49")
    sheet["A4"] = "导出时间"
    sheet["B4"] = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M")
    counts = Counter(row["status"] for row in rows)
    run_count = sum(bool(row.get("qualification_run_id") or row.get("first_run_executed")) for row in rows)
    for address, value in {
        "A5": "用例总数", "B5": len(rows), "C5": "实际运行", "D5": run_count,
        "E5": "未运行", "F5": len(rows) - run_count, "G5": "通过", "H5": counts["PASSED"],
        "A6": "未通过", "B6": counts["FAILED"], "C6": "前置阻塞", "D6": counts["BLOCKED"],
        "E6": "未执行", "F6": counts["NOT_RUN"], "G6": "已运行且有结论",
        "H6": sum(row["status"] in {"PASSED", "FAILED"} and bool(row.get("qualification_run_id") or row.get("first_run_executed")) for row in rows),
    }.items():
        sheet[address] = value
    blocked_without_run = sum(row["status"] == "BLOCKED" and not (row.get("qualification_run_id") or row.get("first_run_executed")) for row in rows)
    sheet["A7"] = f"未运行的 {len(rows) - run_count} 条包括 {blocked_without_run} 条前置阻塞和 {counts['NOT_RUN']} 条未执行。"
    sheet["A7"].font = Font(name="微软雅黑", size=10, italic=True, color="64748B")
    for row in sheet["A5:I6"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="EEF3F9")
            cell.font = Font(name="微软雅黑", size=10, bold=cell.column % 2 == 0, color="152C49")

    for column, title in enumerate(HEADERS, 1):
        cell = sheet.cell(9, column, title)
        cell.fill = PatternFill("solid", fgColor="243F67")
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[9].height = 27

    for index, case in enumerate(rows, 10):
        result = _run_result(runs_dir, case.get("qualification_run_id"))
        values = (
            case["case_id"], _display_text(case.get("title")), _display_text(case.get("case_type")) or "未填写",
            "已运行" if case.get("qualification_run_id") or case.get("first_run_executed") else "未运行",
            STATUS_LABELS.get(case["status"], case["status"]), _elapsed_seconds(result),
            _display_text(case.get("steps")), _display_text(case.get("expected")), _conclusion(case, result),
        )
        for column, value in enumerate(values, 1):
            if isinstance(value, str):
                _text_cell(sheet, index, column, value)
            else:
                sheet.cell(index, column, value)
            cell = sheet.cell(index, column)
            cell.font = Font(name="微软雅黑", size=10, color="1E293B")
            cell.alignment = Alignment(vertical="center")
            if index % 2:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
        sheet.cell(index, 6).number_format = "0.0"
        state_cell = sheet.cell(index, 5)
        state_color = {"PASSED": "187048", "FAILED": "B91C1C", "BLOCKED": "9A5B00"}.get(case["status"], "64748B")
        state_cell.font = Font(name="微软雅黑", size=10, bold=True, color=state_color)
        sheet.row_dimensions[index].height = 22

    for column, width in enumerate((17, 46, 17, 14, 16, 20, 72, 76, 105), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
