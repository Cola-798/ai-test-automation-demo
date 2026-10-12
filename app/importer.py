from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from .schemas import SourceCase


HEADER_ALIASES = {
    "case_id": ("用例编号", "用例ID", "case_id"),
    "title": ("用例标题", "标题", "title"),
    "precondition": ("前置条件", "precondition"),
    "steps": ("步骤", "操作步骤", "steps"),
    "expected": ("预期", "预期结果", "expected"),
    "priority": ("优先级（P0~P3）", "优先级", "priority"),
    "case_type": ("用例类型", "类型", "case_type"),
}


class WorkbookFormatError(ValueError):
    pass


def import_cases(content: bytes) -> list[SourceCase]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookFormatError(f"无法读取 Excel 文件：{exc}") from exc

    imported: list[SourceCase] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration:
            continue
        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        mapping = _resolve_headers(headers)
        if "case_id" not in mapping or "title" not in mapping:
            continue

        for row_number, values in enumerate(rows, start=2):
            case_id = _cell(values, mapping["case_id"])
            title = _cell(values, mapping["title"])
            if not case_id and not title:
                continue
            if not case_id or not title:
                raise WorkbookFormatError(f"{sheet.title} 第 {row_number} 行缺少用例编号或标题")
            imported.append(
                SourceCase(
                    case_id=case_id,
                    title=title,
                    precondition=_mapped_cell(values, mapping, "precondition"),
                    steps=_mapped_cell(values, mapping, "steps"),
                    expected=_mapped_cell(values, mapping, "expected"),
                    priority=_mapped_cell(values, mapping, "priority"),
                    case_type=_mapped_cell(values, mapping, "case_type"),
                    source_sheet=sheet.title,
                    source_row=row_number,
                )
            )

    if not imported:
        raise WorkbookFormatError("没有找到包含“用例编号”和“用例标题”的有效工作表")
    return imported


def _resolve_headers(headers: list[str]) -> dict[str, int]:
    resolved: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in headers:
                resolved[field] = headers.index(alias)
                break
    return resolved


def _mapped_cell(values: tuple, mapping: dict[str, int], field: str) -> str:
    return _cell(values, mapping[field]) if field in mapping else ""


def _cell(values: tuple, index: int) -> str:
    if index >= len(values) or values[index] is None:
        return ""
    return str(values[index]).strip()
